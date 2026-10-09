from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from bloodsmear.batch.aggregation import BatchSummary
from bloodsmear.batch.domain import BatchJobDetail, BatchMode, JobStatus
from bloodsmear.domain import InferenceResult, ImageInfo, ModelInfo, QCInfo, ResultSummary, SampleMetadata, TimingInfo
from bloodsmear.reporting.context import BatchReportContext
from bloodsmear.reporting.excel import ExcelReportRenderer
from bloodsmear.reporting.pdf import PdfReportRenderer


NOW = datetime(2026, 10, 5, tzinfo=timezone.utc)


def context(mode: BatchMode) -> BatchReportContext:
    metadata = SampleMetadata(
        patient_id="P001" if mode == BatchMode.GROUPED else None,
        stain_method="瑞氏染色", scanner_model="Scanner-A", magnification="100x",
        pixel_resolution="0.25 um/px", operator="Operator-01", notes="批量备注",
    )
    summary = ResultSummary(total_detected_cells=1, cell_counts={"RBC": 1}, all_cell_ratios={"RBC": 1.0}, wbc_differential_ratios={})
    job = BatchJobDetail(
        id="job", mode=mode, sample_id="S001" if mode == BatchMode.GROUPED else None,
        metadata=metadata, status=JobStatus.COMPLETED, total_items=1, completed_items=1,
        created_at=NOW, input_expires_at=NOW + timedelta(hours=24), result_expires_at=NOW + timedelta(days=7), items=[],
    )
    batch = BatchSummary(
        job_id="job", mode=mode, status="completed", sample_id=job.sample_id,
        total_views=1, successful_views=1, failed_views=0, items=[],
        aggregate_summary=summary if mode == BatchMode.GROUPED else None, warnings=[], failures=[],
    )
    result = InferenceResult(
        sample_id="S001", metadata=metadata, created_at=NOW,
        model=ModelInfo(name="model", version="1", sha256="A" * 64),
        image=ImageInfo(filename="x.png", width=10, height=10, sha256="B" * 64),
        detections=[], summary=summary, qc=QCInfo(status="pass", warnings=[]),
        runtime=TimingInfo(provider="CUDAExecutionProvider", preprocess_ms=1, inference_ms=2, postprocess_ms=1, total_ms=4),
    )
    return BatchReportContext(job=job, summary=batch, results=[result], license_status="pending", generated_at=NOW)


def test_grouped_excel_and_pdf_display_all_metadata(tmp_path: Path) -> None:
    ctx = context(BatchMode.GROUPED)
    xlsx = ExcelReportRenderer().render_batch(ctx, tmp_path / "grouped.xlsx")
    pdf = PdfReportRenderer().render_batch(ctx, tmp_path / "grouped.pdf")
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    rows = list(wb["Sample Summary"].values); wb.close()
    assert ("Patient ID", "P001") in rows
    assert ("染色方式", "瑞氏染色") in rows
    assert ("备注", "批量备注") in rows
    text = "\n".join(page.extract_text() or "" for page in PdfReader(pdf).pages)
    assert "P001" in text and "瑞氏染色" in text and "批量备注" in text


def test_independent_reports_show_common_metadata_without_patient_id(tmp_path: Path) -> None:
    ctx = context(BatchMode.INDEPENDENT)
    xlsx = ExcelReportRenderer().render_batch(ctx, tmp_path / "independent.xlsx")
    pdf = PdfReportRenderer().render_batch(ctx, tmp_path / "independent.pdf")
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    rows = list(wb["Overview"].values); wb.close()
    assert ("染色方式", "瑞氏染色") in rows
    assert not any(row and row[0] == "Patient ID" for row in rows)
    text = "\n".join(page.extract_text() or "" for page in PdfReader(pdf).pages)
    assert "瑞氏染色" in text and "Patient ID" not in text
