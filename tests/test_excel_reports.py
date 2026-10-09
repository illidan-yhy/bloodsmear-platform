from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from bloodsmear.domain import (
    Detection, ImageInfo, InferenceResult, ModelInfo, QCInfo,
    ResultSummary, SampleMetadata, TimingInfo,
)
from bloodsmear.reporting.context import ReportContextFactory
from bloodsmear.reporting.excel import ExcelReportRenderer


def result() -> InferenceResult:
    return InferenceResult(
        sample_id="S001",
        metadata=SampleMetadata(patient_id="P001", stain_method="瑞氏染色"),
        created_at=datetime(2026, 10, 4, tzinfo=timezone.utc),
        model=ModelInfo(name="model", version="1", sha256="A" * 64),
        image=ImageInfo(filename="sample.png", width=100, height=80, sha256="B" * 64),
        detections=[Detection(class_id=6, class_name="RBC", confidence=0.91234, bbox_xyxy=(1, 2, 10, 20))],
        summary=ResultSummary(
            total_detected_cells=1,
            cell_counts={"RBC": 1},
            all_cell_ratios={"RBC": 1.0},
            wbc_differential_ratios={"Neutrophil": None},
        ),
        qc=QCInfo(status="pass", warnings=[]),
        runtime=TimingInfo(provider="CUDAExecutionProvider", preprocess_ms=1, inference_ms=2, postprocess_ms=1, total_ms=4),
    )


def test_single_excel_has_fixed_sheets_and_core_values(tmp_path: Path) -> None:
    destination = tmp_path / "report.xlsx"
    ExcelReportRenderer().render_single(
        ReportContextFactory.single(result(), "pending_publisher_confirmation"),
        destination,
    )
    workbook = load_workbook(destination, read_only=True, data_only=True)
    assert workbook.sheetnames == ["Summary", "Detections", "QC", "Runtime"]
    summary = list(workbook["Summary"].values)
    assert ("Sample ID", "S001") in summary
    assert ("染色方式", "瑞氏染色") in summary
    detections = list(workbook["Detections"].values)
    assert detections[1][2] == "RBC"


def test_excel_generation_replaces_atomically_without_temp_files(tmp_path: Path) -> None:
    destination = tmp_path / "report.xlsx"
    destination.write_bytes(b"old")
    ExcelReportRenderer().render_single(
        ReportContextFactory.single(result(), "pending"), destination
    )
    assert destination.read_bytes()[:2] == b"PK"
    assert list(tmp_path.glob("*.tmp")) == []
