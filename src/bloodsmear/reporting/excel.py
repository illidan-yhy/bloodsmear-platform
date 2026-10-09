from __future__ import annotations

import os
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

from bloodsmear.batch.domain import BatchMode
from bloodsmear.reporting.context import (
    BatchReportContext,
    SingleReportContext,
    excel_safe_text,
    shanghai_time,
)


class ExcelReportRenderer:
    def render_single(self, context: SingleReportContext, destination: Path) -> Path:
        workbook = Workbook(write_only=True)
        self._single_summary(workbook, context)
        self._detections(workbook, context.result.detections, context.result.sample_id)
        self._qc(workbook, context.result.qc.status, context.result.qc.warnings)
        self._runtime(workbook, [context.result], context.generated_at)
        return _atomic_workbook(workbook, Path(destination))

    def render_batch(self, context: BatchReportContext, destination: Path) -> Path:
        workbook = Workbook(write_only=True)
        if context.job.mode == BatchMode.INDEPENDENT:
            overview = workbook.create_sheet("Overview")
            _append_pairs(overview, [
                ("Job ID", context.job.id), ("模式", "independent"),
                ("总样本", context.summary.total_views),
                ("成功", context.summary.successful_views), ("失败", context.summary.failed_views),
            ])
            _append_pairs(overview, _metadata_rows(context.job.metadata, include_patient=False))
            self._sample_rows(workbook.create_sheet("Samples"), context.results)
        else:
            sheet = workbook.create_sheet("Sample Summary")
            _append_pairs(sheet, [("Job ID", context.job.id), ("Sample ID", context.job.sample_id or "—")])
            _append_pairs(sheet, _metadata_rows(context.job.metadata, include_patient=True))
            if context.summary.aggregate_summary:
                _append_summary(sheet, context.summary.aggregate_summary)
            self._sample_rows(workbook.create_sheet("Views"), context.results)
        detections = workbook.create_sheet("Detections")
        _append_header(detections, ["Sample ID", "Class ID", "Class", "Confidence", "x1", "y1", "x2", "y2"])
        for result in context.results:
            for item in result.detections:
                detections.append([result.sample_id, item.class_id, item.class_name, item.confidence, *item.bbox_xyxy])
        failures = workbook.create_sheet("Failures")
        _append_header(failures, ["Item ID", "Filename", "Error Code", "Error Message"])
        for failure in context.summary.failures:
            failures.append([failure.get("item_id"), excel_safe_text(failure.get("original_filename")), failure.get("error_code"), excel_safe_text(failure.get("error_message"))])
        self._runtime(workbook, context.results, context.generated_at)
        return _atomic_workbook(workbook, Path(destination))

    def _single_summary(self, workbook: Workbook, context: SingleReportContext) -> None:
        result = context.result
        sheet = workbook.create_sheet("Summary")
        rows = [("报告类型", "单图推理报告"), ("Sample ID", result.sample_id)]
        if result.metadata:
            labels = {
                "patient_id": "Patient ID", "stain_method": "染色方式", "scanner_model": "扫描仪",
                "magnification": "倍率", "pixel_resolution": "像素分辨率", "operator": "操作员", "notes": "备注",
            }
            rows.extend((labels[name], excel_safe_text(value)) for name, value in result.metadata.model_dump().items() if value is not None)
        rows.extend([
            ("图像", excel_safe_text(result.image.filename)), ("图像 SHA-256", result.image.sha256),
            ("模型", f"{result.model.name} {result.model.version}"), ("模型 SHA-256", result.model.sha256),
            ("许可证状态", context.license_status), ("用途", "科研与内部分析，不用于临床诊断"),
        ])
        _append_pairs(sheet, rows)
        _append_summary(sheet, result.summary)

    def _detections(self, workbook: Workbook, detections, sample_id: str) -> None:
        sheet = workbook.create_sheet("Detections")
        _append_header(sheet, ["Ordinal", "Class ID", "Class", "Confidence", "x1", "y1", "x2", "y2"])
        for index, item in enumerate(detections, 1):
            sheet.append([index, item.class_id, item.class_name, item.confidence, *item.bbox_xyxy])

    def _qc(self, workbook: Workbook, status: str, warnings: list[str]) -> None:
        sheet = workbook.create_sheet("QC")
        _append_pairs(sheet, [("Status", status), ("Warnings", ", ".join(warnings) if warnings else "无")])

    def _runtime(self, workbook: Workbook, results, generated_at) -> None:
        sheet = workbook.create_sheet("Runtime")
        _append_header(sheet, ["Sample ID", "Provider", "Preprocess ms", "Inference ms", "Postprocess ms", "Total ms"])
        for result in results:
            runtime = result.runtime
            sheet.append([result.sample_id, runtime.provider, runtime.preprocess_ms, runtime.inference_ms, runtime.postprocess_ms, runtime.total_ms])
        sheet.append(["Generated UTC", generated_at.isoformat()])
        sheet.append(["Generated Asia/Shanghai", shanghai_time(generated_at).isoformat()])

    def _sample_rows(self, sheet, results) -> None:
        names = list(results[0].summary.cell_counts) if results else []
        _append_header(sheet, ["Sample ID", "Filename", *names, "QC", "Total ms"])
        for result in results:
            sheet.append([result.sample_id, excel_safe_text(result.image.filename), *[result.summary.cell_counts[name] for name in names], result.qc.status, result.runtime.total_ms])


def _append_pairs(sheet, rows) -> None:
    for label, value in rows:
        sheet.append([label, value])


def _metadata_rows(metadata, *, include_patient: bool):
    if metadata is None:
        return []
    labels = {
        "patient_id": "Patient ID", "stain_method": "染色方式", "scanner_model": "扫描仪",
        "magnification": "倍率", "pixel_resolution": "像素分辨率", "operator": "操作员", "notes": "备注",
    }
    rows = []
    for key, value in metadata.model_dump().items():
        if value is None or (key == "patient_id" and not include_patient):
            continue
        rows.append((labels[key], excel_safe_text(value)))
    return rows


def _append_summary(sheet, summary) -> None:
    sheet.append([])
    _append_header(sheet, ["Cell Type", "Count", "All Cell Ratio", "WBC Ratio"])
    for name, count in summary.cell_counts.items():
        sheet.append([name, count, summary.all_cell_ratios.get(name), summary.wbc_differential_ratios.get(name)])


def _append_header(sheet, values) -> None:
    cells = []
    for value in values:
        cell = WriteOnlyCell(sheet, value=value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cells.append(cell)
    sheet.append(cells)


def _atomic_workbook(workbook: Workbook, destination: Path) -> Path:
    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, name = tempfile.mkstemp(prefix=f".{destination.stem}.", suffix=".xlsx", dir=destination.parent)
    os.close(descriptor)
    temporary = Path(name)
    try:
        workbook.save(temporary)
        verified = load_workbook(temporary, read_only=True)
        try:
            if not verified.sheetnames:
                raise ValueError("Workbook has no sheets")
        finally:
            verified.close()
        os.replace(temporary, destination)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    return destination
