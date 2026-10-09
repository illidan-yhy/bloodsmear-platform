from __future__ import annotations

import argparse
import asyncio
import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from batch_smoke_test import run as run_batch
from bloodsmear.config import AppSettings
from bloodsmear.inference import InferenceService
from bloodsmear.reporting.service import ReportService


def verify_pdf(path: Path) -> None:
    reader = PdfReader(path)
    assert reader.pages
    for page in reader.pages:
        xobjects = page["/Resources"].get("/XObject", {})
        assert all(obj.get_object().get("/Subtype") != "/Image" for obj in xobjects.values())


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    settings = AppSettings(output_dir=args.output)
    service = InferenceService.from_settings(settings)
    data = Path("samples/Blood.png").read_bytes()
    result = service.infer_bytes(data, "Blood.png", "Blood")
    single = ReportService().generate_single(
        result, service.package.manifest.license_status, args.output / "single"
    )
    workbook = load_workbook(single.excel, read_only=True)
    assert workbook.sheetnames == ["Summary", "Detections", "QC", "Runtime"]
    workbook.close()
    verify_pdf(single.pdf)
    batch_payload = asyncio.run(run_batch(args.output / "batch"))
    with zipfile.ZipFile(batch_payload["zip"]) as archive:
        assert "report.xlsx" in archive.namelist()
        assert "report.pdf" in archive.namelist()
    print(json.dumps({"single_excel": str(single.excel), "single_pdf": str(single.pdf), "batch": batch_payload}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
